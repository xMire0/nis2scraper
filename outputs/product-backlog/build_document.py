from pathlib import Path
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[2]
source = (ROOT / 'docs/PRODUCT-BACKLOG.md').read_text()
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = sec.bottom_margin = Cm(1.8)
sec.left_margin = sec.right_margin = Cm(2)
sec.footer_distance = Cm(.8)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet', 'List Number']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(7)
    st.paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(32)
doc.styles['Title'].font.bold = True
for st in doc.styles:
    for border in st.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
for name, size in [('Heading 1', 21), ('Heading 2', 15), ('Heading 3', 12)]:
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].font.bold = True
    doc.styles[name].paragraph_format.space_before = Pt(15)
    doc.styles[name].paragraph_format.keep_with_next = True
doc.core_properties.title = 'Northstar produktbacklog'
doc.core_properties.subject = 'Prioriterede funktioner og leveranceplan'
doc.core_properties.author = 'Northstar'
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Northstar  |  Produktbacklog  |  ').font.size = Pt(9)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)

def para(text, style=None):
    p = doc.add_paragraph(style=style)
    for part in re.split(r'(\*\*.*?\*\*|`[^`]+`)', text):
        r = p.add_run(part.strip('*`') if part.startswith(('**', '`')) else part)
        if part.startswith('**'): r.bold = True
    return p

def heading(text, level=1):
    text = re.sub(r'[^\w\sæøåÆØÅ]', ' ', text)
    return doc.add_heading(re.sub(r'\s+', ' ', text).strip(), level)

def table(rows, widths):
    t = doc.add_table(rows=1, cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, w in zip(t.columns, widths): col.width = Cm(w)
    for i, row in enumerate(rows):
        cells = t.rows[0].cells if i == 0 else t.add_row().cells
        for j, (cell, text) in enumerate(zip(cells, row)):
            cell.width = Cm(widths[j]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2); p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.line_spacing = 1.03
            r = p.add_run(text); r.font.size = Pt(10)
            if i == 0: r.bold = True; r.font.color.rgb = RGBColor(255,255,255)
            pr = cell._tc.get_or_add_tcPr()
            sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), '17354A' if i == 0 else ('F1F5F7' if i % 2 else 'FFFFFF')); pr.append(sh)
            borders = OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e = OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
            pr.append(borders)
            margins = OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge); e.set(qn('w:w'),'85'); e.set(qn('w:type'),'dxa'); margins.append(e)
            pr.append(margins)
        trpr=t.rows[i]._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
        if i == 0: trpr.append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

title_p = doc.add_paragraph('Northstar produktbacklog', 'Title')
border = OxmlElement('w:pBdr')
for side in ['top', 'left', 'bottom', 'right', 'between']:
    el = OxmlElement('w:' + side); el.set(qn('w:val'), 'nil'); border.append(el)
title_p._p.get_or_add_pPr().append(border)
doc.add_paragraph('Prioriterede funktioner og leveranceplan', 'Subtitle')
para('23. september 2026  •  20 opgaver  •  Tre leverancer')
heading('Anbefalet næste skridt')
para('Prioritér korrekte hjemmesider, tydelig datakvalitet og kontrol over køen. Når vi kan stole på de berigede leads, bygger vi videre på ICP, kontakter og salgsmateriale til det daglige salgsarbejde.')
para('Dokumentet samler de eksisterende funktioner, de vigtigste mangler og konkrete acceptkriterier. Backloggen er et prioriteret forslag; opgaverne er ikke kalenderestimeret eller meldt færdige.')
table([['Leverance','Fokus','Opgaver'], ['A','Pålidelige leads og kontrolleret berigelse','PB-01–05'], ['B','Dagligt salgsarbejde og brugbart materiale','PB-06–10 samt enkel eksport'], ['C','Aktualitet, forbrug og salgsresultater','PB-12–14 og behovsstyrede udvidelser']], [2.5,9,5.5])
heading('Produktmål', 2)
product = source.split('## Produktmål\n\n')[1].split('\n## Hvad findes allerede?')[0]
for p in product.split('\n\n'):
    if p.startswith('Backloggen bygger'): continue
    if p.startswith('Den eksisterende'): continue
    para(p)
para('**Vurderingsgrundlag:** Funktionerne er gennemgået i kode og dokumentation. En repræsentativ live-pilot er endnu ikke gennemført; den er en del af næste leverance.')

rest=source.split('## Hvad findes allerede?\n',1)[1]
code = rest.split('## Kodegrundlag\n', 1)[1]
rest = rest.split('## Kodegrundlag\n', 1)[0]
rest = rest.replace('## Foreslået leveranceplan', '## Kodegrundlag\n' + code + '\n## Foreslået leveranceplan')
lines=('## Hvad findes allerede\n'+rest).splitlines()
i=0
while i < len(lines):
    line=lines[i].strip(); i+=1
    if not line: continue
    if line.startswith('## '):
        title=line[3:]
        hp = heading(title)
        if title.startswith(('Hvad findes','P0:', 'P1:', 'P2:', 'Foreslået leveranceplan')):
            hp.paragraph_format.page_break_before = True
    elif line.startswith('### '):
        title=line[4:]
        if title.startswith('PB-'):
            bits=title.split(' · ')
            heading(bits[1],2)
            p=para(f'{bits[0]}  |  P0  |  Størrelse {bits[-1]}'); p.paragraph_format.keep_with_next=True
        else: heading(title,2)
    elif line.startswith('|'):
        rows=[line]
        while i<len(lines) and lines[i].strip().startswith('|'):
            rows.append(lines[i].strip()); i+=1
        parsed=[[c.strip() for c in r.strip('|').split('|')] for r in rows if not re.match(r'^\|[-| :]+\|$',r)]
        if parsed[0][0]=='ID':
            priority='P1' if parsed[1][0]=='PB-06' else 'P2'
            for row in parsed[1:]:
                ident, title, gap, acceptance=row
                name, size=title.rsplit(' · ',1)
                start = len(doc.paragraphs)
                heading(name,2)
                p=para(f'{ident}  |  {priority}  |  Størrelse {size}'); p.paragraph_format.keep_with_next=True
                para(('**Nuværende hul:** ' if priority=='P1' else '**Hvornår:** ')+gap)
                para('**Acceptkriterier:** '+acceptance)
                for p in doc.paragraphs[start:-1]: p.paragraph_format.keep_with_next = True
        else: table(parsed,[3.0,6.7,7.3])
    elif line.startswith('- '): para(line[2:],'List Bullet')
    elif re.match(r'^\d+\. ',line): para(re.sub(r'^\d+\. ','',line),'List Number')
    else: para(line)

# Keep each P0 task together so its criteria never start without its heading.
ps = doc.paragraphs
for idx, p in enumerate(ps):
    if re.match(r'PB-0[1-5]  \|', p.text):
        end = idx + 1
        while end < len(ps) and not ps[end].style.name.startswith('Heading') and not ps[end]._p.xpath('.//w:br'):
            end += 1
        for item in ps[idx:end-1]: item.paragraph_format.keep_with_next = True
out=ROOT/'outputs/product-backlog/Northstar produktbacklog.docx'
doc.save(out)
print(out)
